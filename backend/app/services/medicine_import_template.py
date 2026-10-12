"""Builds the downloadable medicines Excel template (Instructions, Medicines, Stock, Example)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.medicine_import_parser import (
    MED_COLUMNS, STOCK_COLUMNS, FORMS, GST_CHOICES, SCHEDULE_CHOICES, BILLING_CHOICES, YN_CHOICES,
    MAX_MEDICINE_ROWS, MAX_STOCK_ROWS,
)

_REQUIRED_FILL = PatternFill("solid", fgColor="FFE3B3")
_OPTIONAL_FILL = PatternFill("solid", fgColor="E7EDF3")
_TITLE_FONT = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")

INSTRUCTIONS = [
    ("HOW TO FILL THIS FILE", None),
    ("Fill the 'Medicines' sheet (one row per product you sell). Optionally fill the 'Stock' sheet with what is on your shelves today. Look at the 'Example' sheet first.", None),
    ("Do not rename the sheets or change the header row. Columns with a star (*) must be filled. Use the dropdown arrows where they appear.", None),
    ("", None),
    ("MEDICINES SHEET", None),
    ("Generic name*", "The salt or common name, e.g. Paracetamol."),
    ("Brand name", "The brand you sell, e.g. Crocin. Leave blank if you sell it only by its generic name. If you sell 2 brands of the same medicine, use 2 rows."),
    ("Strength*", "e.g. 500 mg, 125 mg/5 ml, 1 g."),
    ("Form*", "Pick from the dropdown (Tablet, Capsule, Syrup, Injection...)."),
    ("Units per pack*", "How many units in one pack. A strip of 15 tablets = 15. A bottle or a vial = 1."),
    ("Billing mode*", "Per pack = the patient is billed whole packs. Per unit = billed per tablet/unit (strips can be cut)."),
    ("MRP per pack*", "The price printed on the pack, in rupees, e.g. 30.50."),
    ("GST %*", "Pick 0, 5, 12, 18 or 28. Write 12 for 12 percent."),
    ("Schedule*", "OTC, H, H1 or X."),
    ("HSN code", "Optional. Usually 4 to 8 digits."),
    ("NPPA ceiling price per unit", "Optional. The government ceiling price for ONE tablet/unit (not for the whole pack)."),
    ("Low-stock alert at", "Optional. Alert when stock falls to this many units. Left blank = 25."),
    ("Category", "Optional, e.g. Antibiotic."),
    ("Active", "Y or N. Left blank = Y."),
    ("", None),
    ("STOCK SHEET (optional)", None),
    ("Each row is ONE batch. Fill Generic name, Brand name, Strength and Form exactly like on the Medicines sheet.", None),
    ("Quantity in packs*", "Number of PACKS (strips/bottles), not tablets. 20 strips of 15 = write 20. We convert to units for you."),
    ("Expiry date*", "Write as dd-mm-yyyy, e.g. 31-12-2027."),
    ("Batch number*", "As printed on the pack."),
    ("", None),
    ("COMMON MISTAKES", None),
    ("1. Typing tablets instead of packs in the Stock sheet.   2. Writing the date as 12/31/2027 (use dd-mm-yyyy).   3. Writing GST as 0.12 (write 12).", None),
    ("4. Same medicine written twice (same generic, brand, strength and form).   5. Different GST or units per pack on 2 brands of the same generic, strength and form.", None),
    ("If something is wrong, the upload shows every problem with its sheet and row number, and NOTHING is saved until the file is clean.", None),
]

EXAMPLE_MEDICINES = [
    ["Paracetamol", "Crocin", "500 mg", "Tablet", 15, "Per pack", 30.5, 12, "OTC", "30049099", None, 25, "Analgesic", "Y"],
    ["Paracetamol", "Dolo", "500 mg", "Tablet", 15, "Per pack", 32, 12, "OTC", "30049099", None, 25, "Analgesic", "Y"],
    ["Azithromycin", None, "500 mg", "Tablet", 3, "Per pack", 120, 12, "H", "30042019", None, 20, "Antibiotic", "Y"],
    ["Amoxicillin", "Mox", "250 mg", "Capsule", 10, "Per unit", 85, 12, "H", None, None, None, "Antibiotic", "Y"],
    ["Ceftriaxone", None, "1 g", "Injection", 1, "Per unit", 95, 12, "H", None, None, 10, "Antibiotic", "Y"],
    ["Cetirizine", "Cetzine", "5 mg/5 ml", "Syrup", 1, "Per pack", 62, 12, "OTC", None, None, None, "Antiallergic", "Y"],
]
EXAMPLE_STOCK = [
    ["Paracetamol", "Crocin", "500 mg", "Tablet", "CR2411", "31-12-2027", 20, "01-10-2026"],
    ["Paracetamol", "Dolo", "500 mg", "Tablet", "DL7788", "30-06-2028", 12, None],
    ["Azithromycin", None, "500 mg", "Tablet", "AZ1190", "31-03-2027", 8, None],
    ["Ceftriaxone", None, "1 g", "Injection", "CF5521", "28-02-2027", 30, None],
]


def _list_dv(items):
    return DataValidation(type="list", formula1='"' + ",".join(str(i) for i in items) + '"',
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Pick from the list", error="Please choose a value from the dropdown list.")


def _header(ws, columns):
    for c, (_, label, required) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = _BOLD
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(30, len(label) + 4))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _col_letter(columns, key):
    return get_column_letter([k for k, _, _ in columns].index(key) + 1)


def _text_format(ws, columns, keys, last_row):
    for key in keys:
        col = [k for k, _, _ in columns].index(key) + 1
        for r in range(2, last_row + 1):
            ws.cell(row=r, column=col).number_format = "@"


def build_medicine_template() -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    for i, (a, b) in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=i, column=1, value=a).alignment = _WRAP
        if b is None and a and a.isupper() or (b is None and a in ("STOCK SHEET (optional)",)):
            ws.cell(row=i, column=1).font = _TITLE_FONT if i == 1 else _BOLD
        if b is not None:
            ws.cell(row=i, column=1).font = _BOLD
            ws.cell(row=i, column=2, value=b).alignment = _WRAP
        elif a:
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)

    # --- Medicines ---
    wm = wb.create_sheet("Medicines")
    _header(wm, MED_COLUMNS)
    last = MAX_MEDICINE_ROWS + 1
    _text_format(wm, MED_COLUMNS, ["generic", "brand", "strength", "hsn", "category"], last)
    for key, items in (("form", FORMS), ("billing", BILLING_CHOICES), ("gst", GST_CHOICES),
                       ("schedule", SCHEDULE_CHOICES), ("active", YN_CHOICES)):
        dv = _list_dv(items)
        wm.add_data_validation(dv)
        L = _col_letter(MED_COLUMNS, key)
        dv.add(f"{L}2:{L}{last}")
    for key, kind in (("pack", "whole"), ("threshold", "whole"), ("price", "decimal"), ("nppa", "decimal")):
        dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Number needed", error="Please enter a number (0 or more).")
        wm.add_data_validation(dv)
        L = _col_letter(MED_COLUMNS, key)
        dv.add(f"{L}2:{L}{last}")

    # --- Stock ---
    wsk = wb.create_sheet("Stock")
    _header(wsk, STOCK_COLUMNS)
    slast = MAX_STOCK_ROWS + 1
    _text_format(wsk, STOCK_COLUMNS, ["generic", "brand", "strength", "batch"], 2001)
    dv = _list_dv(FORMS)
    wsk.add_data_validation(dv)
    L = _col_letter(STOCK_COLUMNS, "form")
    dv.add(f"{L}2:{L}{slast}")
    dvq = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                         showErrorMessage=True, errorTitle="Number needed", error="Enter the number of PACKS (more than 0).")
    wsk.add_data_validation(dvq)
    L = _col_letter(STOCK_COLUMNS, "qty")
    dvq.add(f"{L}2:{L}{slast}")
    for key in ("expiry", "received"):
        L = _col_letter(STOCK_COLUMNS, key)
        for r in range(2, 2002):
            wsk[f"{L}{r}"].number_format = "dd-mm-yyyy"

    # --- Example ---
    we = wb.create_sheet("Example")
    we["A1"] = "EXAMPLE - Medicines sheet (do not upload this sheet's rows; copy the style)"
    we["A1"].font = _TITLE_FONT
    for c, (_, label, required) in enumerate(MED_COLUMNS, start=1):
        cell = we.cell(row=3, column=c, value=label)
        cell.font = _BOLD
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
        we.column_dimensions[get_column_letter(c)].width = max(14, min(30, len(label) + 4))
    for r, row in enumerate(EXAMPLE_MEDICINES, start=4):
        for c, val in enumerate(row, start=1):
            we.cell(row=r, column=c, value=val)
    start = 4 + len(EXAMPLE_MEDICINES) + 2
    we.cell(row=start, column=1, value="EXAMPLE - Stock sheet (Quantity is in PACKS: 20 strips of 15 tablets = 20)").font = _TITLE_FONT
    for c, (_, label, required) in enumerate(STOCK_COLUMNS, start=1):
        cell = we.cell(row=start + 2, column=c, value=label)
        cell.font = _BOLD
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
    for r, row in enumerate(EXAMPLE_STOCK, start=start + 3):
        for c, val in enumerate(row, start=1):
            we.cell(row=r, column=c, value=val)

    wb.active = 1   # open on the Medicines sheet
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()