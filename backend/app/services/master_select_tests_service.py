"""Hospital side of the MASTER test list: choose tests (ticks or a two-column Excel file), set the price,
and the hospital gets its OWN copy (same fields the manual Add Test screen writes, plus master_test_id for
traceability). PLAN (read-only) -> PREVIEW (stored on the server) -> APPLY (one transaction) -> UNDO.
Import-batch records use kind "master_tests" so the existing Excel-import history/undo is never touched."""
import io
import json
import math
import secrets
from datetime import timedelta
from difflib import get_close_matches
from typing import Optional

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models.import_batch import ImportBatch, ImportBatchItem, ImportPreview
from app.models.master_test import MasterTest, MasterTestParameter
from app.models.test_catalog import TestCatalogItem
from app.models.test_catalog_parameter import TestCatalogParameter
from app.models.test_order import TestOrder
from app.services.import_common import Issue, ImportFileError, is_blank, norm_text, parse_number, read_table
from app.services.master_catalog_service import is_excluded_test
from app.services.medicine_import_service import ImportRefused
from app.utils.audit import stage_action
from app.utils.master_catalog import norm_name
from app.utils.timezone import now_ist_naive

KIND = "master_tests"
PREVIEW_MINUTES = 30
MAX_SELECT_ROWS = 2000
MAX_TICKED = 500
MAX_PRICE = 1_000_000          # same upper limit as the manual Add Test screen
SUGGESTION_ROWS_MAX = 200
COLUMNS = [("name", "Test name*", True), ("price", "Price*", True)]


# ======================= helpers =======================

def clean_price(value):
    """Returns (price, None) or (None, reason). Must be greater than 0 and within the existing limit."""
    n, err = parse_number(value)
    if err:
        return None, err
    if n is None:
        return None, "missing"
    if n <= 0:
        return None, "Price must be greater than 0"
    if n > MAX_PRICE:
        return None, "Price must not be more than 10,00,000"
    return round(n, 2), None


def _param_counts(db: Session, ids):
    from sqlalchemy import func
    if not ids:
        return {}
    return dict(db.query(MasterTestParameter.master_test_id, func.count(MasterTestParameter.id)).filter(
        MasterTestParameter.is_active == True,  # noqa: E712
        MasterTestParameter.master_test_id.in_(list(ids))).group_by(MasterTestParameter.master_test_id).all())


def _param_names(db: Session, ids):
    """{master_test_id: [parameter names in display order]} for the expandable panel rows."""
    out = {}
    if not ids:
        return out
    for mid, name in db.query(MasterTestParameter.master_test_id, MasterTestParameter.name).filter(
            MasterTestParameter.is_active == True,  # noqa: E712
            MasterTestParameter.master_test_id.in_(list(ids))
    ).order_by(MasterTestParameter.display_order, MasterTestParameter.id).all():
        out.setdefault(mid, []).append(name)
    return out


def active_masters(db: Session):
    return [m for m in db.query(MasterTest).filter(MasterTest.is_active == True).all()  # noqa: E712
            if not is_excluded_test(m.name, m.category)]


def hospital_name_state(db: Session, hospital_id: int) -> dict:
    """{normalized name: is_active_flag_of_any_row}. Any existing test with the name counts as 'already added'."""
    state = {}
    for name, active in db.query(TestCatalogItem.name, TestCatalogItem.is_active).filter(
            TestCatalogItem.hospital_id == hospital_id).all():
        k = norm_name(name)
        state[k] = bool(state.get(k)) or bool(active)
    return state


# ======================= list for the two-sided panel =======================

def list_for_hospital(db: Session, hospital_id: int, search: Optional[str], category: Optional[str]) -> dict:
    have = hospital_name_state(db, hospital_id)
    masters = active_masters(db)
    categories = sorted({m.category for m in masters if m.category}, key=str.lower)
    s = norm_name(search) if search else ""
    rows = []
    for m in masters:
        if category and (m.category or "") != category:
            continue
        if s and s not in norm_name(m.name) and s not in norm_name(m.aliases):
            continue
        rows.append(m)
    rows.sort(key=lambda m: m.name_norm)
    counts = _param_counts(db, [m.id for m in rows])
    names = _param_names(db, [m.id for m in rows if m.is_panel])
    items = []
    for m in rows:
        k = norm_name(m.name)
        items.append({"id": m.id, "test_name": m.name, "category": m.category or "", "is_panel": bool(m.is_panel),
                      "parameter_count": counts.get(m.id, 0) if m.is_panel else 0,
                      "parameters": names.get(m.id, []) if m.is_panel else [],
                      "already_added": k in have, "already_inactive": (k in have and not have[k])})
    return {"items": items, "total": len(items), "categories": categories}


# ======================= copy on select (the write) =======================

def _copy_values(m: MasterTest, price: float) -> dict:
    """Exactly what the manual add-test flow writes (create_test in routers/tests.py)."""
    p = bool(m.is_panel)
    return dict(
        name=(m.name or "").strip(), fee=price, category=(m.category or "").strip(),
        purpose=(m.purpose or "").strip(), is_panel=p,
        reference_range_male="" if p else (m.reference_range_male or "").strip(),
        reference_range_female="" if p else (m.reference_range_female or "").strip(),
        unit="" if p else (m.unit or "").strip(), turnaround_hours=m.turnaround_hours,
        aliases=(m.aliases or "").strip(),
        critical_low=None if p else m.critical_low, critical_high=None if p else m.critical_high,
        fasting_required=bool(m.fasting_required), required_tube=(m.required_tube or "").strip() or None,
        is_irreplaceable_sample=False, is_nabl_accredited=False, is_hiv_test=False,
        notifiable_disease_id=None, is_active=True, master_test_id=m.id)


def apply_selection(db: Session, actor, hospital, selections: list, file_name: str, source: str) -> dict:
    """selections = [(MasterTest, price)]. ONE transaction: tests, parameters, batch record, audit row."""
    hid = hospital.id
    params_by_master = {}
    for p in db.query(MasterTestParameter).filter(
            MasterTestParameter.master_test_id.in_([m.id for m, _ in selections]),
            MasterTestParameter.is_active == True  # noqa: E712
    ).order_by(MasterTestParameter.display_order, MasterTestParameter.id).all():
        params_by_master.setdefault(p.master_test_id, []).append(p)
    for m, _ in selections:
        if m.is_panel and not params_by_master.get(m.id):
            raise HTTPException(status_code=409, detail=f"'{m.name}' has no parameters in the master list, so it cannot be added. Please tell support.")
    summary = {"tests_added": len(selections), "panels": sum(1 for m, _ in selections if m.is_panel), "source": source}
    batch = ImportBatch(hospital_id=hid, kind=KIND, created_by=actor.id,
                        created_by_name=(f"{actor.title} {actor.name}" if getattr(actor, "title", None) else actor.name),
                        created_by_role=actor.role.value, file_name=(file_name or "")[:200],
                        status="applied", summary=json.dumps(summary))
    db.add(batch)
    db.flush()
    for m, price in selections:
        t = TestCatalogItem(hospital_id=hid, **_copy_values(m, price))
        db.add(t)
        db.flush()
        db.add(ImportBatchItem(batch_id=batch.id, entity="test", entity_id=t.id, action="created"))
        if m.is_panel:
            for i, p in enumerate(params_by_master.get(m.id, [])):
                tp = TestCatalogParameter(
                    test_catalog_item_id=t.id, hospital_id=hid, name=(p.name or "").strip(),
                    unit=(p.unit or "").strip(), reference_range_male=(p.reference_range_male or "").strip(),
                    reference_range_female=(p.reference_range_female or "").strip(), purpose=(p.purpose or "").strip(),
                    display_order=i, is_active=True, critical_low=p.critical_low, critical_high=p.critical_high)
                db.add(tp)
                db.flush()
                db.add(ImportBatchItem(batch_id=batch.id, entity="test_parameter", entity_id=tp.id, action="created"))
    by = " by super admin" if actor.role.value == "super_admin" else ""
    stage_action(db, actor, "master_tests_added", "import_batch", batch.id, hospital.name,
                 details=f"tests_added={summary['tests_added']}, panels={summary['panels']}, source={source}{by}",
                 hospital_id=hid)
    db.commit()
    n = len(selections)
    return {"import_id": batch.id, "added": n, "message": f"{n} test{'s' if n != 1 else ''} added"}


# ======================= add by ticks =======================

def add_selected(db: Session, actor, hospital, items: list) -> dict:
    if not items:
        raise HTTPException(status_code=400, detail="Tick at least one test.")
    if len(items) > MAX_TICKED:
        raise HTTPException(status_code=400, detail=f"Please add at most {MAX_TICKED} tests at a time.")
    ids = [i.master_test_id for i in items]
    if len(set(ids)) != len(ids):
        raise HTTPException(status_code=400, detail="The same test was selected twice.")
    masters = {m.id: m for m in db.query(MasterTest).filter(MasterTest.id.in_(ids)).all()}
    have = hospital_name_state(db, hospital.id)
    selections, missing_price, bad, existing = [], [], [], []
    for it in items:
        m = masters.get(it.master_test_id)
        if m is None or not m.is_active or is_excluded_test(m.name, m.category):
            raise HTTPException(status_code=404, detail="One of the selected tests is no longer in the master list. Please reload the list.")
        if norm_name(m.name) in have:
            existing.append(m.name)
            continue
        price, err = clean_price(it.price)
        if err == "missing":
            missing_price.append(m.name)
        elif err:
            bad.append(f"{m.name}: {err}")
        else:
            selections.append((m, price))
    if existing:
        raise HTTPException(status_code=409, detail="Already added: " + ", ".join(existing[:5]) + (" and more" if len(existing) > 5 else "") + ". Nothing was saved. Please reload the list.")
    if missing_price:
        raise HTTPException(status_code=400, detail="Enter a price for: " + ", ".join(missing_price[:5]) + (" and more" if len(missing_price) > 5 else ""))
    if bad:
        raise HTTPException(status_code=400, detail="; ".join(bad[:3]))
    return apply_selection(db, actor, hospital, selections, "Selected from master list", "ticks")


# ======================= Excel: template =======================

def build_select_template(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    masters = sorted(active_masters(db), key=lambda m: m.name_norm)
    counts = _param_counts(db, [m.id for m in masters])
    wb = Workbook()
    ws = wb.active
    ws.title = "Tests"
    bold = Font(bold=True)
    req = PatternFill("solid", fgColor="FFE3B3")
    for c, (_, label, _r) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font, cell.fill = bold, req
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 16
    ws.freeze_panes = "A2"
    for r in range(2, MAX_SELECT_ROWS + 2):
        ws.cell(row=r, column=1).number_format = "@"
    wm = wb.create_sheet("Master list")
    for c, h in enumerate(("Test name", "Category", "Type"), start=1):
        wm.cell(row=1, column=c, value=h).font = bold
    wm.column_dimensions["A"].width = 48
    wm.column_dimensions["B"].width = 24
    wm.column_dimensions["C"].width = 24
    for r, m in enumerate(masters, start=2):
        for c, v in ((1, m.name), (2, m.category or ""),
                     (3, f"Panel - {counts.get(m.id, 0)} parameters" if m.is_panel else "Single")):
            cell = wm.cell(row=r, column=c, value=v)
            cell.data_type = "s"       # never let a name that starts with '=' become a formula
    last_master = max(len(masters) + 1, 2)
    dv = DataValidation(type="list", formula1=f"='Master list'!$A$2:$A${last_master}", allow_blank=True,
                        showErrorMessage=True, errorStyle="warning", errorTitle="Not in the master list",
                        error="This name is not in the master list. Pick a name from the dropdown.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{MAX_SELECT_ROWS + 1}")
    dp = DataValidation(type="decimal", operator="between", formula1="0.01", formula2=str(MAX_PRICE), allow_blank=True,
                        showErrorMessage=True, errorTitle="Price needed", error="Enter a price greater than 0.")
    ws.add_data_validation(dp)
    dp.add(f"B2:B{MAX_SELECT_ROWS + 1}")
    wi = wb.create_sheet("Instructions")
    wi.column_dimensions["A"].width = 110
    lines = [
        "HOW TO FILL THIS FILE",
        "Use the 'Tests' sheet. Two columns only: Test name and Price.",
        "Test name: pick it from the dropdown. It must be exactly as in the 'Master list' sheet.",
        "Price: your price in rupees. It must be more than 0.",
        "One row per test. Do not repeat a test.",
        "Tests you already have are shown as 'Already added' and are not changed.",
        "After you upload, you will see a preview. Nothing is saved until you press Confirm.",
        "If any row has a problem (red), nothing is saved until it is fixed.",
    ]
    for i, t in enumerate(lines, start=1):
        c = wi.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=14)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================= Excel: parse + plan =======================

def parse_select_workbook(content: bytes):
    """Returns (rows, issues). rows = [{"row", "name", "price", "forced"}]; raw price kept as text/number."""
    import openpyxl
    issues = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Could not read this file. Make sure it is a valid, unprotected .xlsx.")
    try:
        table = read_table(wb, "Tests", COLUMNS, MAX_SELECT_ROWS, issues, True)
    finally:
        wb.close()
    if table is not None and not table:
        issues.append(Issue("Tests", None, None, "The Tests sheet has no rows."))
    rows = []
    for n, v in table or []:
        price = v["price"]
        if isinstance(price, float) and math.isnan(price):
            price = None
        rows.append({"row": n, "name": norm_text(v["name"]),
                     "price": None if is_blank(price) else (price if isinstance(price, (int, float)) and not isinstance(price, bool) else norm_text(price)),
                     "forced": None})
    return rows, issues


def _suggest(name_key: str, pool: dict):
    """pool = {normalized name: MasterTest}. Up to 3 most similar names (simple, dependency-free)."""
    if not name_key:
        return []
    return [pool[k] for k in get_close_matches(name_key, list(pool.keys()), n=3, cutoff=0.5)]


def plan_rows(db: Session, hospital_id: int, rows: list) -> dict:
    masters = active_masters(db)
    by_key = {norm_name(m.name): m for m in masters}
    by_id = {m.id: m for m in masters}
    have = hospital_name_state(db, hospital_id)
    counts_p = _param_counts(db, [m.id for m in masters if m.is_panel])
    out, seen_master, seen_key, sugg_budget = [], {}, {}, SUGGESTION_ROWS_MAX
    for r in rows:
        name, key = r["name"], norm_name(r["name"])
        row = {"row": r["row"], "name": name, "status": "ok", "message": "", "price": None, "master_test_id": None,
               "master_name": "", "is_panel": False, "parameter_count": 0, "suggestions": []}
        price, perr = clean_price(r["price"])
        m = by_id.get(r.get("forced")) if r.get("forced") else by_key.get(key)
        if not name:
            row.update(status="not_found", message="Test name is empty.")
        elif m is None and is_excluded_test(name):
            row.update(status="ignored", message="HIV and radiology tests are not part of the master list. Ignored.")
        elif m is None:
            row.update(status="not_found", message="Not found in the master list.")
            if sugg_budget > 0:
                sugg_budget -= 1
                row["suggestions"] = [{"master_test_id": s.id, "name": s.name} for s in _suggest(key, by_key)]
        else:
            row.update(master_test_id=m.id, master_name=m.name, is_panel=bool(m.is_panel),
                       parameter_count=counts_p.get(m.id, 0) if m.is_panel else 0)
            mk = norm_name(m.name)
            if mk in have:
                row.update(status="already", message="Already added. It will not be changed.")
            elif mk in seen_master:
                row.update(status="duplicate", message=f"Duplicate of row {seen_master[mk]}. Keep one row.")
            elif perr == "missing":
                row.update(status="price_missing", message="Price is missing.")
            elif perr:
                row.update(status="bad_price", message=perr)
            else:
                row["price"] = price
                seen_master[mk] = r["row"]
        if row["status"] == "not_found" and key and key in seen_key:
            row.update(status="duplicate", message=f"Duplicate of row {seen_key[key]}. Keep one row.", suggestions=[])
        elif key:
            seen_key.setdefault(key, r["row"])
        out.append(row)
    return {"rows": out}


BLOCKING = ("not_found", "duplicate", "bad_price", "price_missing")


def summarize(rows: list) -> dict:
    c = {s: sum(1 for r in rows if r["status"] == s) for s in ("ok", "already", "not_found", "price_missing", "bad_price", "duplicate", "ignored")}
    parts = [f"{c['ok']} will be added", f"{c['already']} already added", f"{c['not_found']} not found"]
    if c["price_missing"]:
        parts.append(f"{c['price_missing']} price missing")
    if c["bad_price"]:
        parts.append(f"{c['bad_price']} bad price")
    if c["duplicate"]:
        parts.append(f"{c['duplicate']} duplicate in file")
    if c["ignored"]:
        parts.append(f"{c['ignored']} ignored")
    return {"counts": c, "text": ", ".join(parts),
            "blocking": sum(c[s] for s in BLOCKING), "will_add": c["ok"]}


def preview_dict(db: Session, hospital_id: int, pv_id, file_name, rows: list, issues: list) -> dict:
    plan = plan_rows(db, hospital_id, rows)
    s = summarize(plan["rows"])
    errors = [i.to_dict() for i in issues]
    warnings = [{"text": f"Row {r['row']}: {r['message']}"} for r in plan["rows"] if r["status"] == "ignored"]
    return {"ok": not errors and s["blocking"] == 0, "preview_id": pv_id, "file_name": file_name,
            "expires_in_minutes": PREVIEW_MINUTES, "summary_text": s["text"], "counts": s["counts"],
            "blocking": s["blocking"], "will_add": s["will_add"], "rows": plan["rows"], "errors": errors,
            "warnings": warnings, "can_confirm": (not errors) and s["blocking"] == 0 and s["will_add"] > 0,
            "nothing_to_add": (not errors) and s["blocking"] == 0 and s["will_add"] == 0}


def build_preview(db: Session, doctor, hospital_id: int, file_name: str, content: bytes) -> dict:
    rows, issues = parse_select_workbook(content)
    if issues or not rows:
        return preview_dict(db, hospital_id, None, file_name, rows, issues)
    now = now_ist_naive()
    db.query(ImportPreview).filter(ImportPreview.expires_at < now).delete(synchronize_session=False)
    pid = secrets.token_urlsafe(24)
    db.add(ImportPreview(id=pid, hospital_id=hospital_id, doctor_id=doctor.id, kind=KIND,
                         file_name=(file_name or "")[:200], payload=json.dumps(rows),
                         created_at=now, expires_at=now + timedelta(minutes=PREVIEW_MINUTES)))
    db.commit()
    return preview_dict(db, hospital_id, pid, file_name, rows, [])


def load_preview(db: Session, doctor, hospital_id: int, preview_id: str) -> ImportPreview:
    pv = db.query(ImportPreview).filter(ImportPreview.id == preview_id).first()
    if pv is None or pv.kind != KIND or pv.doctor_id != doctor.id or pv.hospital_id != hospital_id:
        raise HTTPException(status_code=404, detail="Preview not found. Please upload the file again.")
    if pv.consumed:
        raise HTTPException(status_code=400, detail="This preview was already used. Upload the file again if you need to repeat it.")
    if pv.expires_at < now_ist_naive():
        raise HTTPException(status_code=400, detail="This preview expired. Please upload the file again.")
    return pv


def accept_suggestion(db: Session, doctor, hospital_id: int, preview_id: str, row_no: int, master_test_id: int) -> dict:
    """The user clicked a 'Did you mean ...?' suggestion. The server re-derives the allowed suggestions."""
    pv = load_preview(db, doctor, hospital_id, preview_id)
    rows = json.loads(pv.payload)
    target = next((r for r in rows if r["row"] == row_no), None)
    if target is None:
        raise HTTPException(status_code=404, detail="That row is not in this preview.")
    masters = {norm_name(m.name): m for m in active_masters(db)}
    allowed = {m.id for m in _suggest(norm_name(target["name"]), masters)}
    if master_test_id not in allowed:
        raise HTTPException(status_code=400, detail="That test is not one of the suggestions for this row.")
    target["forced"] = master_test_id
    pv.payload = json.dumps(rows)
    db.commit()
    return preview_dict(db, hospital_id, pv.id, pv.file_name, rows, [])


def confirm_preview(db: Session, doctor, hospital, preview_id: str) -> dict:
    pv = load_preview(db, doctor, hospital.id, preview_id)
    rows = json.loads(pv.payload)
    plan = plan_rows(db, hospital.id, rows)["rows"]       # re-checked against the catalog as it is NOW
    s = summarize(plan)
    if s["blocking"]:
        raise HTTPException(status_code=409, detail={
            "message": "Some rows still have problems (or the catalog changed since the preview). Nothing was saved.",
            "errors": [{"text": f"Row {r['row']}: {r['message']}"} for r in plan if r["status"] in BLOCKING][:200]})
    if not s["will_add"]:
        raise HTTPException(status_code=400, detail="Nothing to add: every test in this file is already added.")
    masters = {m.id: m for m in db.query(MasterTest).filter(MasterTest.id.in_([r["master_test_id"] for r in plan if r["status"] == "ok"])).all()}
    selections = [(masters[r["master_test_id"]], r["price"]) for r in plan if r["status"] == "ok"]
    claimed = db.query(ImportPreview).filter(ImportPreview.id == pv.id, ImportPreview.consumed == False).update(  # noqa: E712
        {"consumed": True}, synchronize_session=False)
    if claimed != 1:
        db.rollback()
        raise HTTPException(status_code=400, detail="This preview was already used.")
    return apply_selection(db, doctor, hospital, selections, pv.file_name, "excel")


# ======================= undo =======================

def undo_selection(db: Session, batch: ImportBatch, actor, hospital) -> dict:
    if batch.status != "applied":
        raise ImportRefused("This was already undone.")
    items = db.query(ImportBatchItem).filter(ImportBatchItem.batch_id == batch.id).all()
    test_ids = [i.entity_id for i in items if i.entity == "test"]
    param_ids = [i.entity_id for i in items if i.entity == "test_parameter"]
    tests = db.query(TestCatalogItem).filter(TestCatalogItem.id.in_(test_ids or [0]),
                                             TestCatalogItem.hospital_id == batch.hospital_id).all()
    problems = []
    for t in tests:
        used = db.query(TestOrder.id).filter(TestOrder.test_id == t.id).count()
        if used:
            problems.append(f"{t.name} has already been ordered ({used} order{'s' if used != 1 else ''}).")
    if problems:
        raise ImportRefused("This cannot be undone because some of these tests are already in use. Nothing was changed.", problems[:10])
    db.query(TestCatalogParameter).filter(TestCatalogParameter.id.in_(param_ids or [0]),
                                          TestCatalogParameter.hospital_id == batch.hospital_id
                                          ).delete(synchronize_session=False)
    db.flush()
    for t in tests:
        db.delete(t)
    db.flush()
    batch.status = "undone"
    batch.undone_at = now_ist_naive()
    batch.undone_by = actor.id
    by = " by super admin" if actor.role.value == "super_admin" else ""
    stage_action(db, actor, "master_tests_undone", "import_batch", batch.id, hospital.name,
                 details=f"tests_removed={len(tests)}{by}", hospital_id=batch.hospital_id)
    db.commit()
    return {"tests_removed": len(tests), "message": f"{len(tests)} test{'s' if len(tests) != 1 else ''} removed"}


def history(db: Session, hospital_id: int) -> list:
    rows = db.query(ImportBatch).filter(ImportBatch.hospital_id == hospital_id, ImportBatch.kind == KIND) \
             .order_by(ImportBatch.id.desc()).limit(10).all()
    out = []
    for b in rows:
        try:
            summ = json.loads(b.summary or "{}")
        except ValueError:
            summ = {}
        out.append({"id": b.id, "file_name": b.file_name, "status": b.status, "tests_added": summ.get("tests_added", 0),
                    "created_at": b.created_at.isoformat() if b.created_at else None, "created_by": b.created_by_name})
    return out