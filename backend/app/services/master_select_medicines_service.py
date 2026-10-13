"""Hospital side of the MASTER medicine list: choose medicines (ticks or a four-column Excel file), set the
MRP per pack, and the hospital gets its OWN generic (parent) row, written exactly the way the manual
Add Medicine screen writes it (stock 0), plus one child brand row per brand entered, written the way the manual
Add Brand screen writes it. master_medicine_id is stored on the parent row for traceability only.
PLAN (read-only) -> PREVIEW (stored on the server) -> APPLY (one transaction) -> UNDO.
Import-batch records use kind "master_medicines" so the existing Excel-import history/undo is never touched."""
import io
import json
import secrets
from datetime import timedelta
from difflib import get_close_matches
from typing import Optional

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.models.admission import AdmissionMedicationOrder
from app.models.hospital_medicine import HospitalMedicine
from app.models.import_batch import ImportBatch, ImportBatchItem, ImportPreview
from app.models.master_medicine import MasterMedicine
from app.models.medicine_batch import MedicineBatch
from app.models.medicine_order import MedicineOrder
from app.routers.medicines import (
    MAX_PACK_SIZE, MAX_PRICE, VALID_BILLING_MODES, VALID_SCHEDULES,
    _ceiling_price_warning, _find_duplicate_medicine, _number_problem, compute_unit_price,
)
from app.services.import_common import (
    ImportFileError, Issue, is_blank, norm_key, norm_strength_key, norm_text, parse_number, parse_whole, read_table,
)
from app.services.medicine_import_service import ImportRefused
from app.utils.audit import stage_action
from app.utils.master_catalog import clean_text, norm_name
from app.utils.timezone import now_ist_naive

KIND = "master_medicines"
PREVIEW_MINUTES = 30
MAX_SELECT_ROWS = 3000
MAX_TICKED = 500
MAX_BRANDS_PER_MEDICINE = 10
MAX_BRAND_LEN = 100
MAX_HSN_LEN = 20
SUGGESTION_ROWS_MAX = 200
DEFAULT_THRESHOLD = 25
COLUMNS = [("name", "Medicine*", True), ("mrp", "MRP per pack*", True),
           ("pack", "Units per pack", False), ("brand", "Brand name", False)]


# ======================= helpers =======================

def clean_mrp(value):
    """Returns (mrp, None) or (None, reason). Must be greater than 0 and within the existing limit."""
    n, err = parse_number(value)
    if err:
        return None, err
    if n is None:
        return None, "missing"
    if n <= 0:
        return None, "MRP must be greater than 0"
    if n > MAX_PRICE:
        return None, f"MRP must not be more than {MAX_PRICE:,}"
    return round(n, 2), None


def clean_pack(value, default: int):
    """Blank means the master default. Returns (pack, None) or (None, reason)."""
    if is_blank(value):
        return int(default or 1), None
    n, err = parse_whole(value)
    if err:
        return None, err
    if n is None or not (1 <= n <= MAX_PACK_SIZE):
        return None, f"Units per pack must be between 1 and {MAX_PACK_SIZE:,}"
    return n, None


def active_masters(db: Session):
    return db.query(MasterMedicine).filter(MasterMedicine.is_active == True).all()  # noqa: E712


def med_key(generic, strength, form):
    return (norm_key(generic), norm_strength_key(strength), norm_key(form))


def hospital_keys(db: Session, hospital_id: int) -> set:
    """Generic (parent) rows the hospital already has (active ones, same rule as the duplicate check)."""
    rows = db.query(HospitalMedicine.generic_name, HospitalMedicine.strength, HospitalMedicine.dosage_forms).filter(
        HospitalMedicine.hospital_id == hospital_id, HospitalMedicine.parent_medicine_id == None,  # noqa: E711
        HospitalMedicine.is_active == True).all()  # noqa: E712
    return {med_key(g, s, f) for g, s, f in rows}


def master_key(m: MasterMedicine):
    return med_key(m.generic_name, m.strength, m.dosage_form)


# ======================= list for the two-sided panel =======================

def list_for_hospital(db: Session, hospital_id: int, search: Optional[str], schedule: Optional[str],
                      form: Optional[str]) -> dict:
    have = hospital_keys(db, hospital_id)
    masters = active_masters(db)
    schedules = sorted({(m.schedule or "").lower() for m in masters if m.schedule})
    forms = sorted({m.dosage_form for m in masters if m.dosage_form}, key=str.lower)
    s = norm_name(search) if search else ""
    sched = (schedule or "").strip().lower()
    rows = []
    for m in masters:
        if sched and (m.schedule or "").lower() != sched:
            continue
        if form and (m.dosage_form or "") != form:
            continue
        if s and s not in m.name_norm and s not in norm_name(m.category) and s not in norm_name(m.generic_name):
            continue
        rows.append(m)
    rows.sort(key=lambda m: m.name_norm)
    items = [{"id": m.id, "display_name": m.display_name, "generic_name": m.generic_name, "strength": m.strength or "",
              "dosage_form": m.dosage_form or "", "schedule": (m.schedule or "otc").lower(),
              "default_pack_size": m.default_pack_size or 1, "default_billing_mode": m.default_billing_mode,
              "gst_percent": m.gst_percent, "category": m.category or "",
              "already_added": master_key(m) in have} for m in rows]
    return {"items": items, "total": len(items), "schedules": schedules, "forms": forms}


# ======================= copy on select (the write) =======================

def _check_master(m: MasterMedicine, pack: int):
    if (m.schedule or "").lower() not in VALID_SCHEDULES:
        raise HTTPException(status_code=409, detail=f"'{m.display_name}' has an invalid schedule in the master list. Please tell support.")
    if (m.default_billing_mode or "") not in VALID_BILLING_MODES:
        raise HTTPException(status_code=409, detail=f"'{m.display_name}' has an invalid billing mode in the master list. Please tell support.")
    problem = _number_problem(gst_percent=m.gst_percent, pack_size=pack)
    if problem:
        raise HTTPException(status_code=409, detail=f"'{m.display_name}': {problem}. Please tell support.")


def apply_selection(db: Session, actor, hospital, entries: list, file_name: str, source: str) -> dict:
    """entries = [{"master", "mrp", "pack", "hsn", "nppa", "thr", "brands": [(name, mrp)]}].
    ONE transaction: generic rows, brand rows, batch record, audit row. Nothing is committed on any failure."""
    hid = hospital.id
    for e in entries:
        m = e["master"]
        _check_master(m, e["pack"])
        if _find_duplicate_medicine(db, hid, m.generic_name, m.strength, m.dosage_form):
            raise HTTPException(status_code=409, detail=f"{m.display_name} is already in your catalog. Nothing was saved. Please reload the list.")
    n_brands = sum(len(e["brands"]) for e in entries)
    summary = {"medicines_added": len(entries), "brands_added": n_brands, "source": source}
    batch = ImportBatch(hospital_id=hid, kind=KIND, created_by=actor.id,
                        created_by_name=(f"{actor.title} {actor.name}" if getattr(actor, "title", None) else actor.name),
                        created_by_role=actor.role.value, file_name=(file_name or "")[:200],
                        status="applied", summary=json.dumps(summary))
    db.add(batch)
    db.flush()
    warnings = []
    for e in entries:
        m = e["master"]
        thr = e["thr"] if e["thr"] is not None else DEFAULT_THRESHOLD
        parent = HospitalMedicine(
            hospital_id=hid, generic_name=(m.generic_name or "").strip(), brand_names="",
            category=(m.category or "").strip(), dosage_forms=(m.dosage_form or "").strip(),
            strength=(m.strength or "").strip(), schedule=(m.schedule or "otc").lower(),
            low_stock_threshold=thr, pack_size=e["pack"], price_per_pack=e["mrp"],
            nppa_ceiling_price=e["nppa"], billing_mode=m.default_billing_mode, gst_percent=m.gst_percent,
            hsn_code=(e["hsn"] or (m.hsn_code or "").strip() or None),
            price=compute_unit_price(e["mrp"], e["pack"]), stock_quantity=0, is_active=True,
            master_medicine_id=m.id)
        db.add(parent)
        db.flush()
        db.add(ImportBatchItem(batch_id=batch.id, entity="medicine", entity_id=parent.id, action="created"))
        w = _ceiling_price_warning(parent)
        if w:
            warnings.append(w)
        for brand_name, brand_mrp in e["brands"]:
            brand = HospitalMedicine(
                hospital_id=hid, generic_name=parent.generic_name, category=parent.category,
                dosage_forms=parent.dosage_forms, strength=parent.strength, schedule=parent.schedule,
                brand_name=brand_name, parent_medicine_id=parent.id, low_stock_threshold=thr,
                pack_size=parent.pack_size, price_per_pack=brand_mrp, billing_mode=parent.billing_mode,
                gst_percent=parent.gst_percent, hsn_code=parent.hsn_code, nppa_ceiling_price=parent.nppa_ceiling_price,
                price=compute_unit_price(brand_mrp, parent.pack_size), stock_quantity=0, is_active=True)
            db.add(brand)
            db.flush()
            db.add(ImportBatchItem(batch_id=batch.id, entity="medicine", entity_id=brand.id, action="created"))
            w = _ceiling_price_warning(brand)
            if w:
                warnings.append(w)
    by = " by super admin" if actor.role.value == "super_admin" else ""
    stage_action(db, actor, "master_medicines_added", "import_batch", batch.id, hospital.name,
                 details=f"medicines_added={len(entries)}, brands_added={n_brands}, source={source}{by}",
                 hospital_id=hid)
    db.commit()
    n = len(entries)
    msg = f"{n} medicine{'s' if n != 1 else ''} added"
    if n_brands:
        msg += f" with {n_brands} brand{'s' if n_brands != 1 else ''}"
    return {"import_id": batch.id, "added": n, "brands_added": n_brands, "message": msg,
            "warnings": list(dict.fromkeys(warnings))[:10]}


# ======================= add by ticks =======================

def _clean_extras(label, hsn, nppa, thr):
    hsn = clean_text(hsn)
    if len(hsn) > MAX_HSN_LEN:
        return None, None, None, f"{label}: HSN code is too long"
    nppa_n, err = parse_number(nppa)
    if err:
        return None, None, None, f"{label}: NPPA ceiling price {err}"
    thr_n, err = parse_whole(thr)
    if err:
        return None, None, None, f"{label}: low-stock alert {err}"
    problem = _number_problem(nppa_ceiling_price=nppa_n, low_stock_threshold=thr_n)
    if problem:
        return None, None, None, f"{label}: {problem}"
    return (hsn or None), (round(nppa_n, 2) if nppa_n is not None else None), thr_n, None


def add_selected(db: Session, actor, hospital, items: list) -> dict:
    if not items:
        raise HTTPException(status_code=400, detail="Tick at least one medicine.")
    if len(items) > MAX_TICKED:
        raise HTTPException(status_code=400, detail=f"Please add at most {MAX_TICKED} medicines at a time.")
    ids = [i.master_medicine_id for i in items]
    if len(set(ids)) != len(ids):
        raise HTTPException(status_code=400, detail="The same medicine was selected twice.")
    masters = {m.id: m for m in db.query(MasterMedicine).filter(MasterMedicine.id.in_(ids)).all()}
    have = hospital_keys(db, hospital.id)
    entries, missing_price, bad, existing = [], [], [], []
    for it in items:
        m = masters.get(it.master_medicine_id)
        if m is None or not m.is_active:
            raise HTTPException(status_code=404, detail="One of the selected medicines is no longer in the master list. Please reload the list.")
        if master_key(m) in have:
            existing.append(m.display_name)
            continue
        mrp, merr = clean_mrp(it.price_per_pack)
        pack, perr = clean_pack(it.pack_size, m.default_pack_size)
        hsn, nppa, thr, xerr = _clean_extras(m.display_name, it.hsn_code, it.nppa_ceiling_price, it.low_stock_threshold)
        if merr == "missing":
            missing_price.append(m.display_name)
            continue
        if merr:
            bad.append(f"{m.display_name}: {merr}")
            continue
        if perr:
            bad.append(f"{m.display_name}: {perr}")
            continue
        if xerr:
            bad.append(xerr)
            continue
        brands, seen = [], set()
        if len(it.brands or []) > MAX_BRANDS_PER_MEDICINE:
            bad.append(f"{m.display_name}: at most {MAX_BRANDS_PER_MEDICINE} brands")
            continue
        for b in (it.brands or []):
            name = clean_text(b.brand_name)
            if not name:
                continue            # an empty brand line is simply ignored
            if len(name) > MAX_BRAND_LEN:
                bad.append(f"{m.display_name}: brand name is too long")
                break
            if norm_name(name) in seen:
                bad.append(f"{m.display_name}: brand '{name}' is entered twice")
                break
            seen.add(norm_name(name))
            if is_blank(b.price_per_pack):
                bmrp = mrp
            else:
                bmrp, berr = clean_mrp(b.price_per_pack)
                if berr:
                    bad.append(f"{m.display_name} ({name}): {'MRP is missing' if berr == 'missing' else berr}")
                    break
            brands.append((name, bmrp))
        else:
            entries.append({"master": m, "mrp": mrp, "pack": pack, "hsn": hsn, "nppa": nppa, "thr": thr, "brands": brands})
    if existing:
        raise HTTPException(status_code=409, detail="Already added: " + ", ".join(existing[:5]) + (" and more" if len(existing) > 5 else "") + ". Nothing was saved. Please reload the list.")
    if missing_price:
        raise HTTPException(status_code=400, detail="Enter the MRP per pack for: " + ", ".join(missing_price[:5]) + (" and more" if len(missing_price) > 5 else ""))
    if bad:
        raise HTTPException(status_code=400, detail="; ".join(bad[:3]))
    return apply_selection(db, actor, hospital, entries, "Selected from master list", "ticks")


# ======================= Excel: template =======================

def build_select_template(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    masters = sorted(active_masters(db), key=lambda m: m.name_norm)
    wb = Workbook()
    ws = wb.active
    ws.title = "Medicines"
    bold = Font(bold=True)
    req = PatternFill("solid", fgColor="FFE3B3")
    for c, (_, label, required) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = bold
        if required:
            cell.fill = req
    for col, w in (("A", 52), ("B", 18), ("C", 16), ("D", 28)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    for r in range(2, MAX_SELECT_ROWS + 2):
        ws.cell(row=r, column=1).number_format = "@"
        ws.cell(row=r, column=4).number_format = "@"
    wm = wb.create_sheet("Master list")
    for c, h in enumerate(("Medicine", "Schedule", "Default units per pack", "GST %"), start=1):
        wm.cell(row=1, column=c, value=h).font = bold
    for col, w in (("A", 52), ("B", 12), ("C", 24), ("D", 10)):
        wm.column_dimensions[col].width = w
    for r, m in enumerate(masters, start=2):
        a = wm.cell(row=r, column=1, value=m.display_name)
        a.data_type = "s"          # never let a name that starts with '=' become a formula
        wm.cell(row=r, column=2, value=(m.schedule or "otc").upper())
        wm.cell(row=r, column=3, value=m.default_pack_size or 1)
        wm.cell(row=r, column=4, value=m.gst_percent)
    last_master = max(len(masters) + 1, 2)
    dv = DataValidation(type="list", formula1=f"='Master list'!$A$2:$A${last_master}", allow_blank=True,
                        showErrorMessage=True, errorStyle="warning", errorTitle="Not in the master list",
                        error="This medicine is not in the master list. Pick it from the dropdown.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{MAX_SELECT_ROWS + 1}")
    dp = DataValidation(type="decimal", operator="between", formula1="0.01", formula2=str(MAX_PRICE), allow_blank=True,
                        showErrorMessage=True, errorTitle="MRP needed", error="Enter the MRP per pack, more than 0.")
    ws.add_data_validation(dp)
    dp.add(f"B2:B{MAX_SELECT_ROWS + 1}")
    dk = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAX_PACK_SIZE), allow_blank=True,
                        showErrorMessage=True, errorTitle="Units per pack", error="Enter a whole number, 1 or more.")
    ws.add_data_validation(dk)
    dk.add(f"C2:C{MAX_SELECT_ROWS + 1}")
    wi = wb.create_sheet("Instructions")
    wi.column_dimensions["A"].width = 110
    lines = [
        "HOW TO FILL THIS FILE",
        "Use the 'Medicines' sheet.",
        "Medicine: pick it from the dropdown. It must be exactly as in the 'Master list' sheet.",
        "MRP per pack: the price printed on the pack, in rupees. It must be more than 0.",
        "Units per pack: optional. Leave it blank to use the default shown in the 'Master list' sheet.",
        "Brand name: optional. To add several brands of one medicine, repeat the row with a different brand and MRP.",
        "Every repeated row of a medicine needs a different brand name, and the same units per pack.",
        "Medicines you already have are shown as 'Already added' and are not changed.",
        "Stock is not added here. Add stock afterwards with 'Add batch' (or the stock import).",
        "After you upload, you will see a preview. Nothing is saved until you press Confirm.",
        "If any row has a problem (red), nothing is saved until it is fixed.",
    ]
    for i, t in enumerate(lines, start=1):
        c = wi.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=14)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================= Excel: parse + plan =======================

def _raw_number(v):
    if is_blank(v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return None if v != v else v
    return norm_text(v)


def parse_select_workbook(content: bytes):
    """Returns (rows, issues). rows = [{"row", "name", "mrp", "pack", "brand", "forced"}]."""
    import openpyxl
    issues = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Could not read this file. Make sure it is a valid, unprotected .xlsx.")
    try:
        table = read_table(wb, "Medicines", COLUMNS, MAX_SELECT_ROWS, issues, True)
    finally:
        wb.close()
    if table is not None and not table:
        issues.append(Issue("Medicines", None, None, "The Medicines sheet has no rows."))
    rows = []
    for n, v in table or []:
        rows.append({"row": n, "name": norm_text(v["name"]), "mrp": _raw_number(v["mrp"]),
                     "pack": _raw_number(v["pack"]), "brand": norm_text(v["brand"]), "forced": None})
    return rows, issues


def _suggest(name_key: str, pool: dict):
    """pool = {normalized name: MasterMedicine}. Up to 3 most similar names (simple, dependency-free)."""
    if not name_key:
        return []
    return [pool[k] for k in get_close_matches(name_key, list(pool.keys()), n=3, cutoff=0.5)]


def plan_rows(db: Session, hospital_id: int, rows: list) -> dict:
    masters = active_masters(db)
    by_key = {m.name_norm: m for m in masters}
    by_id = {m.id: m for m in masters}
    have = hospital_keys(db, hospital_id)
    out, seen_master, seen_pair, seen_key, sugg_budget = [], {}, {}, {}, SUGGESTION_ROWS_MAX
    for r in rows:
        name, key, brand = r["name"], norm_name(r["name"]), r.get("brand") or ""
        row = {"row": r["row"], "name": name, "brand": brand, "status": "ok", "message": "", "mrp": None,
               "pack": None, "master_medicine_id": None, "master_name": "", "schedule": "", "default_pack": None,
               "suggestions": []}
        m = by_id.get(r.get("forced")) if r.get("forced") else by_key.get(key)
        if not name:
            row.update(status="not_found", message="Medicine name is empty.")
        elif m is None:
            row.update(status="not_found", message="Not found in the master list.")
            if sugg_budget > 0:
                sugg_budget -= 1
                row["suggestions"] = [{"master_medicine_id": s.id, "name": s.display_name} for s in _suggest(key, by_key)]
        else:
            row.update(master_medicine_id=m.id, master_name=m.display_name, schedule=(m.schedule or "otc").upper(),
                       default_pack=m.default_pack_size or 1)
            mrp, merr = clean_mrp(r["mrp"])
            pack, perr = clean_pack(r.get("pack"), m.default_pack_size)
            pair = (m.id, norm_name(brand))
            first = seen_master.get(m.id)
            if master_key(m) in have:
                row.update(status="already", message="Already added. It will not be changed.")
            elif pair in seen_pair:
                row.update(status="duplicate", message=f"Duplicate of row {seen_pair[pair]}. Keep one row.")
            elif first is not None and not brand:
                row.update(status="duplicate", message=f"This medicine is already on row {first['row']}. To add another brand, write a Brand name on this row.")
            elif merr == "missing":
                row.update(status="price_missing", message="MRP per pack is missing.")
            elif merr:
                row.update(status="bad_price", message=merr)
            elif perr:
                row.update(status="bad_pack", message=perr)
            elif len(brand) > MAX_BRAND_LEN:
                row.update(status="bad_pack", message=f"Brand name is too long (maximum {MAX_BRAND_LEN} characters).")
            elif first is not None and pack != first["pack"]:
                row.update(status="bad_pack", message=f"Units per pack must be the same on every row of this medicine (row {first['row']} uses {first['pack']}).")
            else:
                row.update(mrp=mrp, pack=pack)
                if first is None:
                    seen_master[m.id] = {"row": r["row"], "pack": pack}
                seen_pair[pair] = r["row"]
        if row["status"] == "not_found" and key and key in seen_key:
            row.update(status="duplicate", message=f"Duplicate of row {seen_key[key]}. Keep one row.", suggestions=[])
        elif key:
            seen_key.setdefault(key, r["row"])
        out.append(row)
    return {"rows": out}


BLOCKING = ("not_found", "duplicate", "bad_price", "bad_pack", "price_missing")


def summarize(rows: list) -> dict:
    c = {s: sum(1 for r in rows if r["status"] == s)
         for s in ("ok", "already", "not_found", "price_missing", "bad_price", "bad_pack", "duplicate")}
    ok_rows = [r for r in rows if r["status"] == "ok"]
    meds = len({r["master_medicine_id"] for r in ok_rows})
    brands = sum(1 for r in ok_rows if r.get("brand"))
    first = f"{meds} will be added" + (f" (with {brands} brand{'s' if brands != 1 else ''})" if brands else "")
    parts = [first, f"{c['already']} already added", f"{c['not_found']} not found"]
    if c["price_missing"]:
        parts.append(f"{c['price_missing']} MRP missing")
    if c["bad_price"]:
        parts.append(f"{c['bad_price']} bad price")
    if c["bad_pack"]:
        parts.append(f"{c['bad_pack']} bad pack or brand")
    if c["duplicate"]:
        parts.append(f"{c['duplicate']} duplicate in file")
    return {"counts": c, "text": ", ".join(parts), "blocking": sum(c[s] for s in BLOCKING),
            "will_add": meds, "brands": brands}


def preview_dict(db: Session, hospital_id: int, pv_id, file_name, rows: list, issues: list) -> dict:
    plan = plan_rows(db, hospital_id, rows)
    s = summarize(plan["rows"])
    errors = [i.to_dict() for i in issues]
    return {"ok": not errors and s["blocking"] == 0, "preview_id": pv_id, "file_name": file_name,
            "expires_in_minutes": PREVIEW_MINUTES, "summary_text": s["text"], "counts": s["counts"],
            "blocking": s["blocking"], "will_add": s["will_add"], "brands": s["brands"], "rows": plan["rows"],
            "errors": errors, "warnings": [],
            "can_confirm": (not errors) and s["blocking"] == 0 and s["will_add"] > 0,
            "nothing_to_add": (not errors) and s["blocking"] == 0 and s["will_add"] == 0}


def build_preview(db: Session, doctor, hospital_id: int, file_name: str, content: bytes) -> dict:
    rows, issues = parse_select_workbook(content)
    if issues or not rows:
        return preview_dict(db, hospital_id, None, file_name, rows, issues)
    now = now_ist_naive()
    db.query(ImportPreview).filter(ImportPreview.expires_at < now).delete(synchronize_session=False)
    pid = secrets.token_urlsafe(24)
    db.add(ImportPreview(id=pid, hospital_id=hospital_id, doctor_id=doctor.id, kind=KIND,
                         file_name=(file_name or "")[:200], payload=json.dumps(rows),
                         created_at=now, expires_at=now + timedelta(minutes=PREVIEW_MINUTES)))
    db.commit()
    return preview_dict(db, hospital_id, pid, file_name, rows, [])


def load_preview(db: Session, doctor, hospital_id: int, preview_id: str) -> ImportPreview:
    pv = db.query(ImportPreview).filter(ImportPreview.id == preview_id).first()
    if pv is None or pv.kind != KIND or pv.doctor_id != doctor.id or pv.hospital_id != hospital_id:
        raise HTTPException(status_code=404, detail="Preview not found. Please upload the file again.")
    if pv.consumed:
        raise HTTPException(status_code=400, detail="This preview was already used. Upload the file again if you need to repeat it.")
    if pv.expires_at < now_ist_naive():
        raise HTTPException(status_code=400, detail="This preview expired. Please upload the file again.")
    return pv


def accept_suggestion(db: Session, doctor, hospital_id: int, preview_id: str, row_no: int, master_medicine_id: int) -> dict:
    """The user clicked a 'Did you mean ...?' suggestion. The server re-derives the allowed suggestions."""
    pv = load_preview(db, doctor, hospital_id, preview_id)
    rows = json.loads(pv.payload)
    target = next((r for r in rows if r["row"] == row_no), None)
    if target is None:
        raise HTTPException(status_code=404, detail="That row is not in this preview.")
    masters = {m.name_norm: m for m in active_masters(db)}
    allowed = {m.id for m in _suggest(norm_name(target["name"]), masters)}
    if master_medicine_id not in allowed:
        raise HTTPException(status_code=400, detail="That medicine is not one of the suggestions for this row.")
    target["forced"] = master_medicine_id
    pv.payload = json.dumps(rows)
    db.commit()
    return preview_dict(db, hospital_id, pv.id, pv.file_name, rows, [])


def confirm_preview(db: Session, doctor, hospital, preview_id: str) -> dict:
    pv = load_preview(db, doctor, hospital.id, preview_id)
    rows = json.loads(pv.payload)
    plan = plan_rows(db, hospital.id, rows)["rows"]       # re-checked against the catalog as it is NOW
    s = summarize(plan)
    if s["blocking"]:
        raise HTTPException(status_code=409, detail={
            "message": "Some rows still have problems (or the catalog changed since the preview). Nothing was saved.",
            "errors": [{"text": f"Row {r['row']}: {r['message']}"} for r in plan if r["status"] in BLOCKING][:200]})
    if not s["will_add"]:
        raise HTTPException(status_code=400, detail="Nothing to add: every medicine in this file is already added.")
    masters = {m.id: m for m in db.query(MasterMedicine).filter(
        MasterMedicine.id.in_([r["master_medicine_id"] for r in plan if r["status"] == "ok"])).all()}
    ents = {}
    for r in plan:
        if r["status"] != "ok":
            continue
        e = ents.get(r["master_medicine_id"])
        if e is None:
            e = ents[r["master_medicine_id"]] = {"master": masters[r["master_medicine_id"]], "mrp": r["mrp"],
                                                 "pack": r["pack"], "hsn": None, "nppa": None, "thr": None, "brands": []}
        if r["brand"]:
            e["brands"].append((r["brand"], r["mrp"]))
    claimed = db.query(ImportPreview).filter(ImportPreview.id == pv.id, ImportPreview.consumed == False).update(  # noqa: E712
        {"consumed": True}, synchronize_session=False)
    if claimed != 1:
        db.rollback()
        raise HTTPException(status_code=400, detail="This preview was already used.")
    return apply_selection(db, doctor, hospital, list(ents.values()), pv.file_name, "excel")


# ======================= undo =======================

def undo_selection(db: Session, batch: ImportBatch, actor, hospital) -> dict:
    if batch.status != "applied":
        raise ImportRefused("This was already undone.")
    items = db.query(ImportBatchItem).filter(ImportBatchItem.batch_id == batch.id, ImportBatchItem.entity == "medicine").all()
    ids = {i.entity_id for i in items}
    meds = db.query(HospitalMedicine).filter(HospitalMedicine.id.in_(ids or [0]),
                                             HospitalMedicine.hospital_id == batch.hospital_id).all()
    problems = []
    for m in meds:
        name = m.generic_name + (f" - {m.brand_name}" if m.brand_name else "")
        if db.query(MedicineOrder.id).filter(MedicineOrder.catalog_medicine_id == m.id).first() or \
           db.query(AdmissionMedicationOrder.id).filter(AdmissionMedicationOrder.medicine_id == m.id).first():
            problems.append(f"{name} has already been used in an order.")
        if db.query(MedicineBatch.id).filter(MedicineBatch.medicine_id == m.id).first():
            problems.append(f"{name} already has stock added.")
        if m.parent_medicine_id is None and db.query(HospitalMedicine.id).filter(
                HospitalMedicine.parent_medicine_id == m.id, ~HospitalMedicine.id.in_(ids or [0])).first():
            problems.append(f"{name} has brands that were added after this.")
    if problems:
        raise ImportRefused("This cannot be undone because some of these medicines are already in use. Nothing was changed.",
                            list(dict.fromkeys(problems))[:10])
    for m in sorted(meds, key=lambda x: x.parent_medicine_id is None):   # brand rows first, then generic rows
        db.delete(m)
        db.flush()
    batch.status = "undone"
    batch.undone_at = now_ist_naive()
    batch.undone_by = actor.id
    parents = sum(1 for m in meds if m.parent_medicine_id is None)
    by = " by super admin" if actor.role.value == "super_admin" else ""
    stage_action(db, actor, "master_medicines_undone", "import_batch", batch.id, hospital.name,
                 details=f"medicines_removed={parents}, brands_removed={len(meds) - parents}{by}",
                 hospital_id=batch.hospital_id)
    db.commit()
    return {"medicines_removed": parents, "message": f"{parents} medicine{'s' if parents != 1 else ''} removed"}


def history(db: Session, hospital_id: int) -> list:
    rows = db.query(ImportBatch).filter(ImportBatch.hospital_id == hospital_id, ImportBatch.kind == KIND) \
             .order_by(ImportBatch.id.desc()).limit(10).all()
    out = []
    for b in rows:
        try:
            summ = json.loads(b.summary or "{}")
        except ValueError:
            summ = {}
        out.append({"id": b.id, "file_name": b.file_name, "status": b.status,
                    "medicines_added": summ.get("medicines_added", 0), "brands_added": summ.get("brands_added", 0),
                    "created_at": b.created_at.isoformat() if b.created_at else None, "created_by": b.created_by_name})
    return out