"""Builds the downloadable lab tests Excel template (Instructions, Tests, Example)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.test_import_parser import TEST_COLUMNS, MAX_TEST_ROWS

_REQUIRED_FILL = PatternFill("solid", fgColor="FFE3B3")
_OPTIONAL_FILL = PatternFill("solid", fgColor="E7EDF3")
_TITLE = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")

INSTRUCTIONS = [
    ("HOW TO FILL THIS FILE", None),
    ("Use the 'Tests' sheet. One row per test. Panels (like CBC) have one MAIN row plus one row per parameter. Look at the 'Example' sheet first.", None),
    ("Do not rename the sheet or change the header row. Columns with a star (*) are required on a main row.", None),
    ("", None),
    ("MAIN ROW (Parent test is EMPTY)", None),
    ("Test name*", "The test as doctors will order it, e.g. Complete Blood Count or Fasting Blood Sugar."),
    ("Has parameters*", "Y = a panel with several results (CBC, Lipid Profile). N = a single-value test (Fasting Blood Sugar). Same as the 'has parameters' tick box on the Add Test screen."),
    ("Price*", "The price in rupees. For a panel, the price goes on the main row only."),
    ("Category, Sample tube, Fasting required, Turnaround hours, Other names", "Optional. Other names = alternative names, separated by commas (e.g. CBC, Hemogram)."),
    ("Unit, Male/Female low/high, Reference text, Critical low/high", "Only for a single-value test (Has parameters = N). Low/High are plain numbers (no negatives). Use Reference text for results like Negative."),
    ("", None),
    ("PARAMETER ROWS (a panel's results)", None),
    ("Parent test", "Type the main test's name exactly (e.g. Complete Blood Count). Then fill Test name (e.g. Haemoglobin), Unit, ranges, Critical values and Order."),
    ("Order", "Optional display order inside the panel (1, 2, 3...). If blank, the file order is used."),
    ("A panel's price, tube, fasting etc.", "Leave empty on parameter rows. Parameters are not separate tests: they cannot be ordered or billed on their own."),
    ("", None),
    ("COMMON MISTAKES", None),
    ("1. Has parameters = Y but no parameter rows.   2. Parameter row with a Parent test that does not match a main row exactly.   3. Has parameters = N but also parameter rows.", None),
    ("4. The same parameter twice in one panel.   5. Low higher than High.   6. Putting the price on parameter rows.", None),
    ("", None),
    ("EXTRA COLUMNS (all optional)", None),
    ("Purpose", "A short note on why the test is done. On a parameter row, the purpose of that parameter."),
    ("Irreplaceable sample", "Y for samples that can never be re-collected (CSF, biopsy, bone marrow). Main rows only."),
    ("NABL accredited", "Y if this test is NABL accredited at your lab (controls the accreditation mark on the report). Main rows only."),
    ("Active", "Y or N. Blank = no change (new tests are Active). N deactivates an existing test; Y reactivates it. Main rows only."),
    ("If something is wrong, the upload shows every problem with its row number, and NOTHING is saved until the file is clean.", None),
]

H = [c[1] for c in TEST_COLUMNS]
# name, parent, has_params, category, price, tube, fasting, tat, aliases, unit, mlow, mhigh, flow, fhigh, reftext, clow, chigh, order
EXAMPLE_ROWS = [
    ["Complete Blood Count", None, "Y", "Haematology", 350, "EDTA (purple top)", "N", 4, "CBC, Hemogram", None, None, None, None, None, None, None, None, None, "Screens for anaemia, infection and clotting problems", "N", "N", "Y"],
    ["Haemoglobin", "Complete Blood Count", None, None, None, None, None, None, None, "g/dL", 13, 17, 12, 15.5, None, 7, 20, 1, "Oxygen-carrying protein in red cells"],
    ["RBC Count", "Complete Blood Count", None, None, None, None, None, None, None, "million/uL", 4.5, 5.5, 4, 5, None, None, None, 2],
    ["WBC Count", "Complete Blood Count", None, None, None, None, None, None, None, "/uL", 4000, 11000, 4000, 11000, None, 2000, 30000, 3],
    ["Platelet Count", "Complete Blood Count", None, None, None, None, None, None, None, "lakh/uL", 1.5, 4.5, 1.5, 4.5, None, 0.5, 10, 4],
    ["Haematocrit", "Complete Blood Count", None, None, None, None, None, None, None, "%", 40, 50, 36, 46, None, None, None, 5],
    ["MCV", "Complete Blood Count", None, None, None, None, None, None, None, "fL", 80, 100, 80, 100, None, None, None, 6],
    ["Fasting Blood Sugar", None, "N", "Biochemistry", 120, "Fluoride (grey top)", "Y", 2, "FBS, Fasting Glucose", "mg/dL", 70, 100, 70, 100, None, 40, 400, None, "Screens for diabetes", "N", "N", "Y"],
    ["Dengue NS1 Antigen", None, "N", "Serology", 600, "Plain (red top)", "N", 6, None, None, None, None, None, None, "Negative", None, None, None],
]


def _header(ws, row):
    for c, (_, label, required) in enumerate(TEST_COLUMNS, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = _BOLD
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(13, min(28, len(label) + 4))


def build_test_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100
    for i, (a, b) in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=i, column=1, value=a).alignment = _WRAP
        if b is not None:
            ws.cell(row=i, column=1).font = _BOLD
            ws.cell(row=i, column=2, value=b).alignment = _WRAP
        elif a:
            if a.isupper() or a.startswith("MAIN ROW") or a.startswith("PARAMETER ROWS"):
                ws.cell(row=i, column=1).font = _TITLE if i == 1 else _BOLD
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)

    wt = wb.create_sheet("Tests")
    _header(wt, 1)
    wt.row_dimensions[1].height = 32
    wt.freeze_panes = "B2"
    last = MAX_TEST_ROWS + 1
    cols = [k for k, _, _ in TEST_COLUMNS]

    def L(key):
        return get_column_letter(cols.index(key) + 1)

    for key in ("name", "parent", "category", "tube", "aliases", "unit", "reftext", "purpose"):
        c = cols.index(key) + 1
        for r in range(2, 2002):
            wt.cell(row=r, column=c).number_format = "@"
    for key in ("has_params", "fasting", "irreplaceable", "nabl", "active"):
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick Y or N", error="Please choose Y or N.")
        wt.add_data_validation(dv)
        dv.add(f"{L(key)}2:{L(key)}{last}")
    for key in ("price", "mlow", "mhigh", "flow", "fhigh", "clow", "chigh"):
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Number needed", error="Please enter a number (0 or more).")
        wt.add_data_validation(dv)
        dv.add(f"{L(key)}2:{L(key)}{last}")
    for key in ("tat", "order"):
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Whole number needed", error="Please enter a whole number.")
        wt.add_data_validation(dv)
        dv.add(f"{L(key)}2:{L(key)}{last}")

    we = wb.create_sheet("Example")
    we["A1"] = "EXAMPLE - Complete Blood Count is a PANEL (Has parameters = Y) with 6 parameter rows. Fasting Blood Sugar is a SINGLE test (N)."
    we["A1"].font = _TITLE
    _header(we, 3)
    for r, row in enumerate(EXAMPLE_ROWS, start=4):
        for c, val in enumerate(row, start=1):
            we.cell(row=r, column=c, value=val)
    wb.active = 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()